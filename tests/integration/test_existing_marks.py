"""입력에 이미 있는 마킹을 걷어내고 시작하는 단계 시험 (변경안 A6).

입력 파일은 그대로 두고, 분석과 사본 마킹이 '마킹 없는 상태'에서 시작하는지 확인한다.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from lxml import etree

from build_synthetic import PDF_YES, YES_TABLE_CELL, YES_WEIGHT
from helpers import FIXTURES, make_services
from patent_marker.analysis import run_analysis
from patent_marker.existing_marks import clear_pdf_marks, clear_pptx_marks
from patent_marker.export.runner import export_run
from patent_marker.feedback.events import record_page_review
from patent_marker.parsers import parse_document
from patent_marker.runtime import sha256_file

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"


def _review(services, file_name: str, yes_fragments: list[str]) -> None:
    """문서의 모든 검토 단위를 일괄 검토한다. 지정한 문장이 든 구간만 YES."""
    document_id = services.database.query_one(
        "SELECT document_id FROM documents WHERE file_name = ? AND is_current = 1", (file_name,))[0]
    rows = services.database.query(
        "SELECT s.segment_id, s.normalized_text, p.unit FROM segments s JOIN paragraphs p USING (paragraph_id) "
        "WHERE s.document_id = ? ORDER BY s.seq", (document_id,))
    for unit in sorted({row["unit"] for row in rows}):
        members = [row for row in rows if row["unit"] == unit]
        yes = [row["segment_id"] for row in members if any(f in row["normalized_text"] for f in yes_fragments)]
        record_page_review(services.database, document_id=document_id, unit=unit, yes_segment_ids=yes,
                           reviewer_id="reviewer-01", guideline_version="1.0")


def _mark(tmp_path: Path, source: Path, run_id: str, yes_fragments: list[str], services=None):
    """source 하나를 분석·검토·내보내기 하고 (services, 사본 항목)을 돌려준다."""
    services = services or make_services(tmp_path)
    inbox = tmp_path / f"inbox-{run_id}"
    inbox.mkdir()
    target = inbox / source.name
    target.write_bytes(source.read_bytes())
    run_analysis(services, inbox, run_id)
    _review(services, source.name, yes_fragments)
    result = export_run(services.database, services.config, run_id, tmp_path / "outputs" / run_id, ["annotated"])
    return services, target, result["annotated"][0]


def _pptx_marks(path: Path) -> dict[str, int]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    names = [shape.name for slide in presentation.slides for shape in slide.shapes]
    return {
        "slides": len(presentation.slides),
        "tags": sum(name == "PM_MARK_TAG" for name in names),
        "outlines": sum(name == "PM_MARK_OUTLINE" for name in names),
        "summary_shapes": sum(name.startswith("PM_SUMMARY_") for name in names),
        "highlights": sum(len(list(slide._element.iter(A + "highlight"))) for slide in presentation.slides),
    }


# ---------------------------------------------------------------- PPTX: 도구가 만든 표시
def test_marked_pptx_copy_is_parsed_like_the_original(tmp_path, config):
    _services, _source, entry = _mark(tmp_path, FIXTURES / "sample_report.pptx", "run-001", [YES_WEIGHT, YES_TABLE_CELL])
    marked = Path(entry["output"])
    original = parse_document(FIXTURES / "sample_report.pptx", config)
    again = parse_document(marked, config)
    assert [(b.unit, b.kind, b.text) for b in again.blocks] == [(b.unit, b.kind, b.text) for b in original.blocks]
    assert [b.locator["shape_id"] for b in again.blocks] == [b.locator["shape_id"] for b in original.blocks]
    assert again.coverage["units_total"] == original.coverage["units_total"] == 8  # 요약 슬라이드는 세지 않는다
    assert original.coverage["existing_marks"] == {"tool_shapes": 0, "tool_slides": 0, "highlights": 0, "comments": 0}
    assert again.coverage["existing_marks"] == {"tool_shapes": 4, "tool_slides": 1, "highlights": 6, "comments": 0}
    warning = next(w for w in again.warnings if w["code"] == "existing_marks_cleared")
    assert "11건" in warning["message"] and "도구 표시 4" in warning["message"]
    assert not any(w["code"] == "existing_marks_cleared" for w in original.warnings)


def test_remarking_a_marked_copy_starts_clean_and_does_not_accumulate(tmp_path):
    services, _source, first = _mark(tmp_path, FIXTURES / "sample_report.pptx", "run-001", [YES_WEIGHT, YES_TABLE_CELL])
    first_marks = _pptx_marks(Path(first["output"]))
    assert first["cleared"] == {"tool_shapes": 0, "tool_slides": 0, "highlights": 0, "comments": 0}
    assert first_marks == {"slides": 9, "tags": 2, "outlines": 2, "summary_shapes": 2, "highlights": 6}

    marked_input = Path(first["output"])
    before = sha256_file(marked_input)
    # 2차: 마킹 사본을 입력으로 넣고 이번에는 표 행만 후보로 판정한다
    _services, second_source, second = _mark(tmp_path, marked_input, "run-002", [YES_TABLE_CELL], services)
    assert sha256_file(marked_input) == before and sha256_file(second_source) == before  # 입력은 그대로
    assert Path(second["output"]).name == "sample_report.marked.pptx"  # .marked.marked 가 되지 않는다
    assert second["cleared"] == {"tool_shapes": 4, "tool_slides": 1, "highlights": 6, "comments": 0}
    second_marks = _pptx_marks(Path(second["output"]))
    # 낡은 표시(슬라이드 3)는 사라지고 새 표시만 남는다. 요약 슬라이드도 한 장뿐이다.
    assert second_marks == {"slides": 9, "tags": 1, "outlines": 1, "summary_shapes": 2, "highlights": 3}


# ---------------------------------------------------------------- PPTX: 작성자의 형광펜·메모
def _authored_pptx(path: Path) -> None:
    """작성자가 직접 형광펜, 기존 메모, 최신 메모, 빨간 사각형을 넣은 덱."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.util import Inches

    presentation = Presentation(str(FIXTURES / "sample_report.pptx"))
    slide = presentation.slides[1]
    run = slide.placeholders[1].text_frame.paragraphs[0].runs[0]._r
    properties = run.find(A + "rPr")
    if properties is None:
        properties = etree.SubElement(run, A + "rPr")
        run.remove(properties)
        run.insert(0, properties)
    highlight = etree.SubElement(properties, A + "highlight")
    etree.SubElement(highlight, A + "srgbClr", val="FFFF00")
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(5), Inches(2), Inches(1))
    box.name = "Author Rectangle"
    box.fill.background()
    box.line.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    legacy = (f'<p:cmLst xmlns:p="{P_NS}"><p:cm authorId="0" dt="2026-10-01T00:00:00.000" idx="1">'
              '<p:pos x="10" y="10"/><p:text>SECRET-COMMENT-LEGACY</p:text></p:cm></p:cmLst>').encode()
    part = Part(PackURI("/ppt/comments/comment1.xml"),
                "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
                presentation.part.package, legacy)
    slide.part.relate_to(part, f"{R_NS}/comments")

    modern = (f'<p188:cmLst xmlns:p188="{P188}"><p188:cm id="{{1}}" authorId="{{2}}" created="2026-10-01T00:00:00.000">'
              '<p188:txBody>SECRET-COMMENT-MODERN</p188:txBody></p188:cm><p188:cm id="{3}" authorId="{2}" '
              'created="2026-10-02T00:00:00.000"><p188:txBody>second</p188:txBody></p188:cm></p188:cmLst>').encode()
    target = presentation.slides[2]
    part = Part(PackURI("/ppt/comments/modernComment_1.xml"), "application/vnd.ms-powerpoint.comments+xml",
                presentation.part.package, modern)
    relationship_id = target.part.relate_to(part, "http://schemas.microsoft.com/office/2018/10/relationships/comments")
    target._element.append(etree.fromstring(
        f'<p:extLst xmlns:p="{P_NS}" xmlns:r="{R_NS}"><p:ext uri="{{6950BFC3-D8DA-4A85-94F7-54DA5524770B}}">'
        f'<p188:commentRel xmlns:p188="{P188}" r:id="{relationship_id}"/></p:ext></p:extLst>'))
    presentation.save(str(path))


def _package_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "\n".join(name + "\n" + archive.read(name).decode("utf-8", "ignore") for name in archive.namelist()
                         if name.endswith((".xml", ".rels")))


def test_author_highlights_and_comments_are_cleared_in_the_copy_only(tmp_path, config):
    from pptx import Presentation

    source = tmp_path / "authored.pptx"
    _authored_pptx(source)
    source_text = _package_text(source)
    assert "SECRET-COMMENT-LEGACY" in source_text and "SECRET-COMMENT-MODERN" in source_text and "FFFF00" in source_text

    parsed = parse_document(source, config)
    assert parsed.coverage["existing_marks"] == {"tool_shapes": 0, "tool_slides": 0, "highlights": 1, "comments": 3}
    reference = parse_document(FIXTURES / "sample_report.pptx", config)
    assert [b.text for b in parsed.blocks] == [b.text for b in reference.blocks]  # 마킹은 본문 추출에 영향이 없다

    before = sha256_file(source)
    _services, copied, entry = _mark(tmp_path, source, "run-001", [YES_WEIGHT])
    assert sha256_file(source) == before and sha256_file(copied) == before
    assert _package_text(copied) == source_text  # 입력 사본에는 형광펜·메모가 그대로 있다
    assert entry["status"] == "ok" and entry["cleared"] == {"tool_shapes": 0, "tool_slides": 0, "highlights": 1, "comments": 3}

    output = Path(entry["output"])
    text = _package_text(output)
    assert "SECRET-COMMENT" not in text and "ppt/comments/" not in text and "commentRel" not in text
    assert "FFFF00" not in text  # 작성자 형광펜은 사본에 남지 않는다
    presentation = Presentation(str(output))  # 다시 열 수 있는 정상 패키지
    marks = _pptx_marks(output)
    assert marks["tags"] == 1 and marks["highlights"] == 3 and marks["slides"] == 9
    names = [shape.name for shape in presentation.slides[1].shapes]
    assert "Author Rectangle" in names  # 일반 도형은 내용일 수 있어 지우지 않는다
    for slide in presentation.slides:
        assert not [rel for rel in slide.part.rels.values() if "comments" in rel.reltype]


def test_clearing_is_idempotent_and_only_in_memory(tmp_path):
    from pptx import Presentation

    source = tmp_path / "authored.pptx"
    _authored_pptx(source)
    before = sha256_file(source)
    presentation = Presentation(str(source))
    assert clear_pptx_marks(presentation) == {"tool_shapes": 0, "tool_slides": 0, "highlights": 1, "comments": 3}
    assert clear_pptx_marks(presentation) == {"tool_shapes": 0, "tool_slides": 0, "highlights": 0, "comments": 0}
    assert sha256_file(source) == before


# ---------------------------------------------------------------- PDF
def _annotated_pdf(path: Path) -> None:
    """검토자가 메모·형광펜을 붙이고, 문서 자체의 링크가 있는 PDF."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Highlight, Link, Text
    from pypdf.generic import ArrayObject, FloatObject, NameObject, TextStringObject

    writer = PdfWriter(clone_from=PdfReader(str(FIXTURES / "sample_report.pdf")))
    note = Text(rect=(400, 700, 420, 720), text="SECRETREVIEWNOTE12345")
    note[NameObject("/T")] = TextStringObject("reviewer")
    writer.add_annotation(0, note)
    quad = ArrayObject([FloatObject(v) for v in (56, 610, 300, 610, 56, 598, 300, 598)])
    mark = Highlight(rect=(56, 598, 300, 610), quad_points=quad)
    mark[NameObject("/T")] = TextStringObject("reviewer")
    writer.add_annotation(1, mark)
    writer.add_annotation(0, Link(rect=(56, 40, 200, 60), target_page_index=2))
    with path.open("wb") as handle:
        writer.write(handle)


def _pdf_annotations(path: Path) -> list[tuple[int, str, str | None]]:
    from pypdf import PdfReader

    found = []
    for number, page in enumerate(PdfReader(str(path)).pages, start=1):
        for reference in page.get("/Annots") or []:
            annotation = reference.get_object()
            author = annotation.get("/T")
            found.append((number, str(annotation["/Subtype"]), str(author) if author is not None else None))
    return found


def test_pdf_review_annotations_are_cleared_and_links_are_kept(tmp_path, config):
    source = tmp_path / "reviewed.pdf"
    _annotated_pdf(source)
    assert b"SECRETREVIEWNOTE12345" in source.read_bytes()
    parsed = parse_document(source, config)
    assert parsed.coverage["existing_marks"] == {"tool_annotations": 0, "annotations": 2}  # 링크는 세지 않는다
    assert any(w["code"] == "existing_marks_cleared" for w in parsed.warnings)
    assert PDF_YES in [b.text for b in parsed.blocks]

    before = sha256_file(source)
    services, copied, entry = _mark(tmp_path, source, "run-001", [PDF_YES[:25]])
    assert sha256_file(source) == before and sha256_file(copied) == before
    assert entry["status"] == "ok" and entry["cleared"] == {"tool_annotations": 0, "annotations": 2}
    output = Path(entry["output"])
    assert _pdf_annotations(output) == [(1, "/Link", None), (1, "/Highlight", "patent-marker")]
    assert b"SECRETREVIEWNOTE12345" not in output.read_bytes()  # 걷어낸 메모가 파일 안에 남지 않는다

    # 마킹 사본을 다시 넣어도 표시가 쌓이지 않는다
    again = parse_document(output, config)
    assert again.coverage["existing_marks"] == {"tool_annotations": 1, "annotations": 0}
    assert [b.text for b in again.blocks] == [b.text for b in parsed.blocks]
    _services, _copied, second = _mark(tmp_path, output, "run-002", [PDF_YES[:25]], services)
    assert second["cleared"] == {"tool_annotations": 1, "annotations": 0}
    assert Path(second["output"]).name == "reviewed.marked.pdf"
    assert _pdf_annotations(Path(second["output"])) == [(1, "/Link", None), (1, "/Highlight", "patent-marker")]


def test_clear_pdf_marks_keeps_non_marking_annotations(tmp_path):
    from pypdf import PdfReader, PdfWriter

    source = tmp_path / "reviewed.pdf"
    _annotated_pdf(source)
    writer = PdfWriter(clone_from=PdfReader(str(source)))
    assert clear_pdf_marks(writer) == {"tool_annotations": 0, "annotations": 2}
    assert clear_pdf_marks(writer) == {"tool_annotations": 0, "annotations": 0}
    assert [str(ref.get_object()["/Subtype"]) for ref in writer.pages[0]["/Annots"]] == ["/Link"]
    assert "/Annots" not in writer.pages[1]


# ---------------------------------------------------------------- 출력 이름
def test_marked_copies_with_the_same_name_do_not_overwrite_each_other(tmp_path):
    services = make_services(tmp_path)
    inbox = tmp_path / "inbox"
    (inbox / "a").mkdir(parents=True)
    (inbox / "b").mkdir()
    (inbox / "a" / "report.pdf").write_bytes((FIXTURES / "sample_report.pdf").read_bytes())
    other = tmp_path / "other.pdf"
    _annotated_pdf(other)
    (inbox / "b" / "report.pdf").write_bytes(other.read_bytes())
    run_analysis(services, inbox, "run-001")
    for row in services.database.query("SELECT document_id, file_name FROM documents"):
        segment = services.database.query_one(
            "SELECT segment_id FROM segments WHERE document_id = ? AND normalized_text LIKE ?",
            (row["document_id"], PDF_YES[:20] + "%"))[0]
        from patent_marker.feedback.events import record_feedback

        record_feedback(services.database, target_type="segment", target_id=segment, label="YES",
                        reviewer_id="reviewer-01", guideline_version="1.0")
    result = export_run(services.database, services.config, "run-001", tmp_path / "out", ["annotated"])
    names = sorted(Path(entry["output"]).name for entry in result["annotated"] if entry["status"] == "ok")
    assert names == ["report.marked-2.pdf", "report.marked.pdf"]
    assert pytest
