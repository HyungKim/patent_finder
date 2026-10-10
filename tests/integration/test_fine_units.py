"""판정 단위를 잘게 하는 설정(segmentation.unit=fine) 시험: 불릿마다, 문장마다, 표는 셀마다.

기본값(paragraph)의 동작은 바뀌지 않아야 하고, 두 단위는 같은 문서를 서로 다른 문서 기록으로 나눠 가진다.
"""
from __future__ import annotations

import pdfplumber

from build_synthetic import YES_TABLE_CELL, YES_WEIGHT, YES_WEIGHT_CHILD_1
from helpers import FIXTURES, make_services
from patent_marker.analysis import run_analysis
from patent_marker.classifiers.registry import get_active, load_registered_bundle
from patent_marker.export.runner import export_run
from patent_marker.feedback.events import record_feedback
from patent_marker.ingest import ingest_file

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _inbox(tmp_path):
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    for path in FIXTURES.iterdir():
        (inbox / path.name).write_bytes(path.read_bytes())
    return inbox


def _segment(services, prefix: str):
    row = services.database.query_one(
        "SELECT segment_id, normalized_text, paragraph_id FROM segments WHERE normalized_text LIKE ?", (prefix + "%",))
    assert row is not None, prefix
    return row


def _highlighted_runs(pptx_path, slide_index: int) -> list[str]:
    from pptx import Presentation

    slide = list(Presentation(str(pptx_path)).slides)[slide_index]
    return ["".join(node.text or "" for node in run.iter(A + "t"))
            for run in slide._element.iter(A + "r") if next(run.iter(A + "highlight"), None) is not None]


def test_fine_unit_splits_bullets_sentences_and_cells(tmp_path):
    inbox = _inbox(tmp_path)
    coarse = make_services(tmp_path / "coarse")
    fine = make_services(tmp_path / "fine", **{"segmentation.unit": "fine"})
    run_analysis(coarse, inbox, "run-001")
    run_analysis(fine, inbox, "run-001")
    count = lambda services: services.database.query_one("SELECT COUNT(*) FROM segments")[0]  # noqa: E731
    assert count(fine) > count(coarse)
    assert fine.database.query_one("SELECT DISTINCT segmentation_version FROM documents")[0] == "segmentation-2-fine"
    assert coarse.database.query_one("SELECT DISTINCT segmentation_version FROM documents")[0] == "segmentation-1"

    # 상위 불릿과 하위 불릿이 각자 구간이다 (기본 모드에서는 한 묶음)
    parent, child = _segment(fine, YES_WEIGHT[:20]), _segment(fine, YES_WEIGHT_CHILD_1[:20])
    assert parent["paragraph_id"] != child["paragraph_id"] and parent["normalized_text"] == YES_WEIGHT
    grouped = _segment(coarse, YES_WEIGHT[:20])
    assert YES_WEIGHT_CHILD_1 in grouped["normalized_text"]
    # 표는 셀마다 구간이고 열 제목이 앞에 붙는다
    cell = _segment(fine, "처리 방식: " + YES_TABLE_CELL[:10])
    assert cell["normalized_text"] == "처리 방식: " + YES_TABLE_CELL
    assert "cell" in __import__("json").loads(fine.database.query_one(
        "SELECT quality_flags FROM segments WHERE segment_id = ?", (cell["segment_id"],))[0])
    # 한 문단 안의 두 문장이 각자 구간이다
    second = _segment(fine, "모듈별 온도 센서 값을 1초 주기로")
    assert not second["normalized_text"].startswith("왼쪽 단")
    first = fine.database.query_one("SELECT normalized_text FROM segments WHERE paragraph_id = ? ORDER BY seq",
                                    (second["paragraph_id"],))[0]
    assert first.startswith("왼쪽 단 첫 문단이다")


def test_fine_unit_marks_only_the_cell_sentence_and_bullet(tmp_path):
    inbox = _inbox(tmp_path)
    services = make_services(tmp_path, **{"segmentation.unit": "fine"})
    run_analysis(services, inbox, "run-001")
    document = services.database.query("SELECT document_id, file_name FROM documents")
    by_name = {row["file_name"]: row["document_id"] for row in document}
    for prefix in ("처리 방식: " + YES_TABLE_CELL[:10], YES_WEIGHT[:20], "모듈별 온도 센서 값을 1초 주기로"):
        record_feedback(services.database, target_type="segment", target_id=_segment(services, prefix)["segment_id"],
                        label="YES", reviewer_id="reviewer-01", guideline_version="1.0")
    # 시스템 후보는 빼고 사람 판정만 표시되게 나머지 후보를 모두 NO로 둔다
    for row in services.database.query(
            "SELECT s.segment_id FROM segments s JOIN latest_predictions lp ON lp.segment_id = s.segment_id "
            "WHERE lp.final_decision = 'CANDIDATE' AND s.document_id IN (?, ?)",
            (by_name["sample_report.pptx"], by_name["sample_report.pdf"])):
        if services.database.query_one("SELECT 1 FROM resolved_labels WHERE target_id = ?", (row["segment_id"],)) is None:
            record_feedback(services.database, target_type="segment", target_id=row["segment_id"], label="NO",
                            reviewer_id="reviewer-01", guideline_version="1.0")
    result = export_run(services.database, services.config, "run-001", tmp_path / "out", ["annotated"])
    entries = {item["file_name"]: item for item in result["annotated"]}

    pptx_path = entries["sample_report.pptx"]["output"]
    assert _highlighted_runs(pptx_path, 2) == [YES_WEIGHT]  # 상위 불릿만. 하위 불릿은 칠해지지 않는다
    assert _highlighted_runs(pptx_path, 3) == [YES_TABLE_CELL]  # 행이 아니라 셀 하나만
    with pdfplumber.open(entries["sample_report.pdf"]["output"]) as pdf:
        found = [(page, annotation) for page in pdf.pages for annotation in page.annots]
        assert len(found) == 1  # PDF에서는 둘째 문장 하나만 YES
        page, annotation = found[0]
        # 뷰어가 칠하는 것은 줄마다의 quad다. 둘째 문장이 줄 중간에서 시작하므로 quad 단위로 덮인 글자를 본다
        quads, pieces = annotation["data"]["QuadPoints"], []
        for q in range(len(quads) // 8):
            x0, y1, x1, _, _, y0, _, _ = (float(v) for v in quads[q * 8: q * 8 + 8])
            box = (x0 - 0.5, page.height - y1 - 0.5, x1 + 0.5, page.height - y0 + 0.5)
            pieces.append((page.crop(box).extract_text() or "").replace("\n", " "))
        covered = " ".join(pieces)
        assert covered.startswith("모듈별 온도 센서 값을 1초 주기로") and "왼쪽 단 첫 문단이다" not in covered


def test_fine_unit_highlights_part_of_a_run_without_touching_the_rest(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    inbox = tmp_path / "inbox"
    inbox.mkdir()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame
    frame.text = "첫 문장은 배경 설명이다. 둘째 문장은 임계값을 넘으면 출력을 낮춘다."
    frame.paragraphs[0].runs[0].font.bold = True  # 서식이 있는 run 하나
    presentation.save(str(inbox / "split.pptx"))

    services = make_services(tmp_path, **{"segmentation.unit": "fine"})
    run_analysis(services, inbox, "run-001")
    rows = services.database.query("SELECT segment_id, normalized_text FROM segments ORDER BY seq")
    assert [row["normalized_text"] for row in rows] == ["첫 문장은 배경 설명이다.", "둘째 문장은 임계값을 넘으면 출력을 낮춘다."]
    record_feedback(services.database, target_type="segment", target_id=rows[1]["segment_id"], label="YES",
                    reviewer_id="reviewer-01", guideline_version="1.0")
    record_feedback(services.database, target_type="segment", target_id=rows[0]["segment_id"], label="NO",
                    reviewer_id="reviewer-01", guideline_version="1.0")
    result = export_run(services.database, services.config, "run-001", tmp_path / "out", ["annotated"])
    output = result["annotated"][0]["output"]

    copied = list(Presentation(output).slides)[0]
    paragraph = next(p for p in copied._element.iter(A + "p") if "둘째 문장" in "".join(t.text or "" for t in p.iter(A + "t")))
    runs = [("".join(t.text or "" for t in run.iter(A + "t")), next(run.iter(A + "highlight"), None) is not None)
            for run in paragraph.findall(A + "r")]
    assert runs == [("첫 문장은 배경 설명이다. ", False), ("둘째 문장은 임계값을 넘으면 출력을 낮춘다.", True)]
    # 나눈 조각마다 원래 서식(굵게)이 그대로 있다
    assert [run.find(A + "rPr").get("b") for run in paragraph.findall(A + "r")] == ["1", "1"]


def test_changing_the_unit_reingests_documents_and_rebuilds_the_seed(tmp_path):
    inbox = _inbox(tmp_path)
    coarse = make_services(tmp_path)
    run_analysis(coarse, inbox, "run-001")
    before = get_active(coarse.database)["model_version"]
    fine = make_services(tmp_path, **{"segmentation.unit": "fine"})
    result = ingest_file(fine, inbox / "sample_report.pptx")
    assert not result.reused
    rows = fine.database.query("SELECT document_id, is_current, segmentation_version FROM documents "
                               "WHERE file_name = 'sample_report.pptx' ORDER BY ingested_at")
    assert [(row["is_current"], row["segmentation_version"]) for row in rows] == [(0, "segmentation-1"), (1, "segmentation-2-fine")]
    run_analysis(fine, inbox, "run-002")
    active = get_active(fine.database)
    assert active["model_version"] != before  # 단위가 다르면 분류기가 호환되지 않아 seed를 다시 만든다
    assert load_registered_bundle(fine.database, active["model_version"], fine.config.base_dir).segmentation_version == "segmentation-2-fine"
