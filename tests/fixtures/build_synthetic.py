#!/usr/bin/env python3
"""비기밀 합성 시험 자료 생성기. 실제 업무 문서를 쓰지 않는다 (스펙 14.1, 15.2).

    python tests/fixtures/build_synthetic.py            # tests/fixtures/synthetic/ 에 생성

PDF 생성에는 개발용 의존성 reportlab이 필요하다. 생성된 파일은 Git에 포함하므로
운영 PC에서는 다시 만들 필요가 없다.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

OUT_DIR = Path(__file__).resolve().parent / "synthetic"

# ---- 시험에서 기대값으로 쓰는 문장 ----
YES_WEIGHT = "Radar confidence가 임계값보다 낮으면 camera feature의 가중치를 증가시킴"
YES_WEIGHT_CHILD_1 = "임계값은 최근 N 프레임의 confidence 이동 평균으로 갱신"
YES_WEIGHT_CHILD_2 = "신뢰도가 회복되면 이전 가중치로 단계적으로 복귀"
EFFECT_BULLET = "야간 평가에서 오검출이 기존 대비 12% 감소"
NO_SCHEDULE = "4분기: 통합 평가 및 최종 보고 예정"
YES_TABLE_CELL = "입력 프레임을 4x4 블록으로 나누고 블록별 노이즈 분산으로 필터 강도를 선택"
YES_GROUP_TEXT = "Confidence 추정기 출력이 0.4 미만이면 가중치 제어기가 camera branch gain을 1.5배로 높인다"
NOTES_TEXT = "발표자 노트: 출원 검토 일정은 10월 중 확정"
PDF_YES = "셀 온도 편차가 3도를 넘으면 냉각수 유량 밸브의 개도를 편차에 비례해 증가시키고, 편차가 1도 이하로 줄면 이전 개도로 복귀한다."
PDF_HYPHEN = "In this scheme the valve controller applies a feed-forward term computed from the predicted heat load of each module."
PDF_HEADER = "합성 예시 보고서 (비기밀)"
DOCX_YES = "게이트 산화막 증착 전에 챔버 압력을 2단계로 낮추고, 1단계 종료 시점은 잔류 가스 분압의 기울기로 판정한다."


def _png_bytes(width: int = 640, height: int = 400) -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), (232, 232, 228))
    draw = ImageDraw.Draw(image)
    for index in range(6):
        left = 40 + index * 95
        draw.rectangle([left, 60 + (index % 3) * 40, left + 70, 300], outline=(90, 90, 90), width=3)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# ------------------------------------------------------------------ PPTX
def build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "센서 융합 인식 성능 개선 과제"
    slide.placeholders[1].text = "3분기 진행 보고 (합성 예시 자료)"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "추진 일정 및 예산"
    frame = slide.placeholders[1].text_frame
    frame.text = "3분기: 야간 주행 데이터 추가 수집 완료"
    frame.add_paragraph().text = NO_SCHEDULE
    frame.add_paragraph().text = "예산 집행률 72%, 잔여 예산은 평가 장비 임차에 사용"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "제안 방식: 신뢰도 기반 가중치 조절"
    frame = slide.placeholders[1].text_frame
    frame.text = YES_WEIGHT
    for text in (YES_WEIGHT_CHILD_1, YES_WEIGHT_CHILD_2):
        paragraph = frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 1
    frame.add_paragraph().text = EFFECT_BULLET
    slide.notes_slide.notes_text_frame.text = NOTES_TEXT

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "모듈별 처리 방식 비교"
    table = slide.shapes.add_table(4, 3, Inches(0.5), Inches(1.6), Inches(9), Inches(3)).table
    rows = [
        ("모듈", "처리 방식", "비고"),
        ("전처리", YES_TABLE_CELL, "신규"),
        ("추적", "기존 칼만 필터 유지", "변경 없음"),
        ("일정", "10월 2주차 통합", "-"),
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    # 제목 자리표시자 없이 텍스트 상자만 쓰는 슬라이드 + 그룹 도형 + 좌우 배치
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "융합 제어 구조"
    group = slide.shapes.add_group_shape()
    for index, label in enumerate(("Radar 전처리", "Confidence 추정기", "가중치 제어기")):
        box = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6 + index * 3.0), Inches(1.5), Inches(2.4), Inches(0.9))
        box.text_frame.text = label
    # XML 순서를 읽기 순서와 다르게: 오른쪽 상자를 먼저 추가
    right = slide.shapes.add_textbox(Inches(5.2), Inches(3.0), Inches(4.3), Inches(1.5))
    right.text_frame.word_wrap = True
    right.text_frame.text = YES_GROUP_TEXT
    left = slide.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(4.3), Inches(1.5))
    left.text_frame.word_wrap = True
    left.text_frame.text = "야간에는 카메라 인식 신뢰도가 감소하여 기존 고정 가중치 융합의 오검출이 늘어난다."

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(_png_bytes()), Inches(0.5), Inches(0.5), Inches(9), Inches(6.2))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "야간 오검출률 비교"
    data = CategoryChartData()
    data.categories = ["기존", "제안"]
    data.add_series("오검출률(%)", (8.4, 7.4))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(4.5), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "야간 시나리오 오검출률"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "향후 계획"
    frame = slide.placeholders[1].text_frame
    frame.text = "다음 분기에 센서 융합 성능 개선을 추진한다."
    frame.add_paragraph().text = "주간 회의는 매주 화요일 오전에 진행"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(20)

    prs.save(str(path))


# ------------------------------------------------------------------ DOCX
def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("1. 개요", level=1)
    document.add_paragraph("본 문서는 공정 조건 검토 결과를 정리한 합성 예시 자료이다. 실제 업무 내용과 무관하다.")
    document.add_heading("2. 제안 구조", level=1)
    document.add_paragraph(DOCX_YES)
    document.add_paragraph("증착 온도는 웨이퍼 중심과 가장자리의 온도 차가 5도를 넘지 않도록 히터 구역별 출력을 보정한다.", style="List Bullet")
    document.add_paragraph("보정 계수는 직전 10매의 두께 측정값으로 갱신한다.", style="List Bullet")
    table = document.add_table(rows=3, cols=3)
    cells = [
        ("항목", "조건", "비고"),
        ("압력", "1단계 5 Torr에서 2단계 0.8 Torr로 전환하며 전환 시점은 분압 기울기 < 0.02 로 판정", "신규"),
        ("회의", "매주 목요일 공정 회의", "-"),
    ]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("표 다음 문단: 위 조건은 3.5 nm 두께 목표에 대해 확인하였다.")
    document.add_heading("3. 일정", level=1)
    document.add_paragraph("다음 분기에 양산 라인 적용 여부를 검토할 예정이다.")
    document.save(str(path))


# ------------------------------------------------------------------ PDF
def build_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    font = "HYGothic-Medium"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    width, height = A4
    c = canvas.Canvas(str(path), pagesize=A4, pageCompression=1)
    state = {"page": 0}

    def frame() -> None:
        state["page"] += 1
        c.setFont(font, 8)
        c.drawString(56, height - 30, PDF_HEADER)
        c.drawCentredString(width / 2, 28, f"- {state['page']} -")

    def lines(x: float, y: float, rows: list[str], size: float = 10.5, leading: float = 16) -> float:
        c.setFont(font, size)
        for row in rows:
            c.drawString(x, y, row)
            y -= leading
        return y

    # 1쪽: 제목, 줄바꿈된 문단, 하이픈, 불릿
    frame()
    c.setFont(font, 18)
    c.drawString(56, height - 90, "배터리 열관리 제어 로직 검토")
    y = lines(56, height - 130, [
        "셀 온도 편차가 3도를 넘으면 냉각수 유량 밸브의 개도를 편차에 비례해 증가시키고, 편차가",
        "1도 이하로 줄면 이전 개도로 복귀한다.",
    ])
    y = lines(56, y - 14, [
        "In this scheme the valve controller applies a feed-forward term computed from the pre-",
        "dicted heat load of each module.",
    ])
    lines(56, y - 14, [
        "• 다음 분기에 열관리 성능 개선을 추진한다.",
        "• 회의는 매주 화요일에 진행한다.",
    ])
    c.showPage()

    # 2쪽: 전체 폭 제목 + 2단 본문
    frame()
    c.setFont(font, 16)
    c.drawString(56, height - 90, "모듈 간 온도 균일화 방법")
    left = [
        "왼쪽 단 첫 문단이다. 모듈별 온도 센서 값을",
        "1초 주기로 수집하고 최댓값과 최솟값의",
        "차이를 편차로 정의한다.",
        "편차가 기준을 넘으면 유량 분배 밸브를",
        "고온 모듈 쪽으로 더 연다.",
    ]
    right = [
        "오른쪽 단 첫 문단이다. 분배 비율은 모듈",
        "열용량의 역수에 비례하도록 계산한다.",
        "계산된 비율은 5초 이동 평균으로 평활화한",
        "뒤 밸브 구동기에 전달한다.",
    ]
    lines(56, height - 130, left)
    lines(320, height - 130, right)
    c.showPage()

    # 3쪽: 테두리 있는 표 + 문단
    frame()
    c.setFont(font, 16)
    c.drawString(56, height - 90, "제어 조건 요약")
    col_x = [56, 150, 430, 540]
    row_y = [height - 120, height - 150, height - 180, height - 210]
    for x in col_x:
        c.line(x, row_y[0], x, row_y[-1])
    for yy in row_y:
        c.line(col_x[0], yy, col_x[-1], yy)
    table = [
        ("항목", "조건", "비고"),
        ("유량 밸브", "편차 3도 초과 시 개도를 편차에 비례해 증가", "신규"),
        ("일정", "11월 시제품 평가", "-"),
    ]
    c.setFont(font, 10)
    for r, row in enumerate(table):
        for col, value in enumerate(row):
            c.drawString(col_x[col] + 6, row_y[r] - 20, value)
    lines(56, height - 250, ["표의 조건은 합성 예시이며 실제 제품과 무관하다."])
    c.showPage()

    # 4쪽: 이미지뿐인 페이지 (스캔 페이지 모사)
    frame()
    c.drawImage(ImageReader(io.BytesIO(_png_bytes(900, 1100))), 56, 80, width=width - 112, height=height - 180)
    c.showPage()
    c.save()


def build_text(directory: Path) -> None:
    (directory / "sample_notes.txt").write_text(
        "주간 업무 메모 (합성 예시)\n"
        "\n"
        "인버터 스위칭 주파수를 부하 전류 구간별로 다르게 설정하고, 구간 경계에서는\n"
        "히스테리시스 폭 5 A를 두어 주파수 전환이 반복되지 않게 한다.\n"
        "\n"
        "- 다음 주 화요일 설계 검토 회의 예정\n"
        "- 시험 장비 대여 신청 완료\n",
        encoding="utf-8",
    )
    (directory / "sample_notes.md").write_text(
        "# 검사 알고리즘 메모 (합성 예시)\n"
        "\n"
        "## 제안\n"
        "\n"
        "결함 후보 영역의 밝기 분포를 기준 패치와 비교하고, KL divergence가 0.3 이상이면 "
        "2차 고해상도 촬영을 요청한다.\n"
        "\n"
        "| 단계 | 처리 | 비고 |\n"
        "|---|---|---|\n"
        "| 1차 | 저해상도 전수 검사 | 기존 |\n"
        "| 2차 | 후보 영역만 고해상도 재촬영 후 분류기 적용 | 신규 |\n"
        "\n"
        "## 일정\n"
        "\n"
        "- 10월: 데이터 수집\n"
        "- 11월: 현장 시험\n",
        encoding="utf-8",
    )


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_pptx(OUT_DIR / "sample_report.pptx")
    build_docx(OUT_DIR / "sample_report.docx")
    build_pdf(OUT_DIR / "sample_report.pdf")
    build_text(OUT_DIR)
    for path in sorted(OUT_DIR.iterdir()):
        print(f"{path.name:24s} {path.stat().st_size:8d} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
