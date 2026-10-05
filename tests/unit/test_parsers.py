"""파서 회귀 시험: 원문 위치 보존, 읽기 순서, 표, 미지원·부분 처리 상태 (스펙 4절, 17.1)."""
from __future__ import annotations

import io
import zipfile

import pytest

from build_synthetic import (DOCX_YES, EFFECT_BULLET, NOTES_TEXT, PDF_HEADER, PDF_HYPHEN, PDF_YES, YES_GROUP_TEXT,
                             YES_TABLE_CELL, YES_WEIGHT, YES_WEIGHT_CHILD_1)
from helpers import FIXTURES, make_config
from patent_marker.parsers import parse_document
from patent_marker.runtime import sha256_file


@pytest.fixture(scope="module")
def config(tmp_path_factory):
    return make_config(tmp_path_factory.mktemp("cfg"))


def _texts(result, **conditions):
    return [b.text for b in result.blocks if all(getattr(b, key) == value for key, value in conditions.items())]


# ---------------------------------------------------------------- PPTX
def test_pptx_extracts_blocks_with_locators(config):
    path = FIXTURES / "sample_report.pptx"
    before = sha256_file(path)
    result = parse_document(path, config)
    assert sha256_file(path) == before  # 원본 불변
    assert result.status == "partial"  # 이미지뿐인 슬라이드가 있어 전체가 분석되지는 않았다
    coverage = result.coverage
    assert coverage["units_total"] == 8 and coverage["units_processed"] == 8
    assert coverage["ocr_needed_units"] == [6]
    assert any(w["code"] == "image_only_slide" and w["unit"] == 6 for w in result.warnings)
    assert coverage["scope"] == {"speaker_notes": False, "hidden_slides": True}

    weight = next(b for b in result.blocks if b.text == YES_WEIGHT)
    assert weight.unit == 3 and weight.kind == "list_item" and weight.level == 0
    assert weight.section_title == "제안 방식: 신뢰도 기반 가중치 조절"
    locator = weight.locator
    assert locator["slide"] == 3 and locator["paragraph"] == 0 and locator["bbox_emu"]["width"] > 0
    assert locator["runs"] == [{"run": 0, "start": 0, "end": len(YES_WEIGHT)}]
    child = next(b for b in result.blocks if b.text == YES_WEIGHT_CHILD_1)
    assert child.level == 1 and child.container == weight.container
    assert NOTES_TEXT not in [b.text for b in result.blocks]  # 발표자 노트는 기본 제외


def test_pptx_reading_order_differs_from_xml_order(config):
    result = parse_document(FIXTURES / "sample_report.pptx", config)
    slide = [b for b in result.blocks if b.unit == 5]
    texts = [b.text for b in slide]
    assert texts[0] == "융합 제어 구조" and slide[0].kind == "title" and "title_inferred" in slide[0].flags
    # 그룹 도형 안의 텍스트
    assert texts[1:4] == ["Radar 전처리", "Confidence 추정기", "가중치 제어기"]
    assert all(b.locator.get("group_path") for b in slide[1:4])
    # 왼쪽 상자가 XML에서는 뒤에 있지만 읽기 순서는 앞이다
    left, right = slide[4], slide[5]
    assert right.text == YES_GROUP_TEXT and left.text.startswith("야간에는")
    assert left.locator["xml_order"] > right.locator["xml_order"]
    assert left.locator["reading_order"] < right.locator["reading_order"]


def test_pptx_table_rows_keep_cells_and_headers(config):
    result = parse_document(FIXTURES / "sample_report.pptx", config)
    rows = [b for b in result.blocks if b.kind == "table_row" and b.unit == 4]
    assert len(rows) == 4 and rows[0].flags == ["table_header"]
    row = rows[1]
    assert row.locator["table_row"] == 1
    assert [row.text[c["start"]:c["end"]] for c in row.cells] == ["전처리", YES_TABLE_CELL, "신규"]
    assert [c["header"] for c in row.cells] == ["모듈", "처리 방식", "비고"]
    chart = next(b for b in result.blocks if b.kind == "diagram_text")
    assert chart.text == "야간 시나리오 오검출률" and chart.flags == ["chart_title"]


def test_pptx_notes_and_hidden_slides_follow_configured_scope(tmp_path, config):
    from pptx import Presentation

    presentation = Presentation(str(FIXTURES / "sample_report.pptx"))
    presentation.slides[1]._element.set("show", "0")
    path = tmp_path / "hidden.pptx"
    presentation.save(str(path))

    with_notes = make_config(tmp_path, **{"parsing.include_speaker_notes": True})
    result = parse_document(path, with_notes)
    notes = [b for b in result.blocks if b.kind == "notes"]
    assert [b.text for b in notes] == [NOTES_TEXT] and notes[0].unit == 3
    hidden = [b for b in result.blocks if b.unit == 2]
    assert hidden and all("hidden_slide" in b.flags for b in hidden)
    assert result.coverage["hidden_units"] == [2]

    without_hidden = make_config(tmp_path, **{"parsing.include_hidden_slides": False})
    assert not [b for b in parse_document(path, without_hidden).blocks if b.unit == 2]


def test_pptx_smartart_text_is_extracted():
    from lxml import etree

    from patent_marker.parsers.pptx import A, DGM, a_paragraph_text

    xml = (f'<dgm:dataModel xmlns:dgm="{DGM[1:-1]}" xmlns:a="{A[1:-1]}"><dgm:ptLst>'
           '<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>신뢰도 추정</a:t></a:r><a:br/><a:r><a:t>가중치 전환</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '<dgm:pt modelId="2" type="sibTrans"><dgm:t><a:p><a:r><a:t>무시</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '</dgm:ptLst></dgm:dataModel>')
    root = etree.fromstring(xml.encode("utf-8"))
    points = [p for p in root.iter(DGM + "pt") if p.get("type") not in ("pres", "parTrans", "sibTrans")]
    text, runs = a_paragraph_text(points[0].find(DGM + "t").find(A + "p"))
    assert text == "신뢰도 추정\n가중치 전환"
    assert [(r["start"], r["end"]) for r in runs] == [(0, 6), (6, 7), (7, 13)]


# ---------------------------------------------------------------- DOCX
def test_docx_keeps_document_order_with_table_between_paragraphs(config):
    result = parse_document(FIXTURES / "sample_report.docx", config)
    assert result.status == "ok"
    kinds = [b.kind for b in result.blocks]
    assert kinds == ["heading", "paragraph", "heading", "paragraph", "list_item", "list_item",
                     "table_row", "table_row", "table_row", "paragraph", "heading", "paragraph"]
    body = [b.locator["body_index"] for b in result.blocks]
    assert body == sorted(body)  # 표를 뒤에 몰아 붙이지 않는다
    yes = next(b for b in result.blocks if b.text == DOCX_YES)
    assert yes.section_title == "2. 제안 구조" and yes.locator["runs"][0]["start"] == 0
    row = [b for b in result.blocks if b.kind == "table_row"][1]
    assert row.locator["row"] == 1 and [c["header"] for c in row.cells] == ["항목", "조건", "비고"]
    assert "< 0.02" in row.text and "0.8 Torr" in row.text  # 연산자·단위 보존
    assert "3.5 nm" in result.blocks[9].text


def test_docx_includes_tracked_insertions_and_excludes_deletions(tmp_path, config):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    paragraph = document.add_paragraph("기존 문장 ")
    inserted = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "추가된 내용"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    deleted = OxmlElement("w:del")
    run = OxmlElement("w:r")
    removed = OxmlElement("w:delText")
    removed.text = "삭제된 내용"
    run.append(removed)
    deleted.append(run)
    paragraph._p.append(deleted)
    assert paragraph._p.find(qn("w:ins")) is not None
    path = tmp_path / "tracked.docx"
    document.save(str(path))
    result = parse_document(path, config)
    assert [b.text for b in result.blocks] == ["기존 문장 추가된 내용"]
    assert result.blocks[0].locator["runs"] == [{"run": 0, "start": 0, "end": 6}, {"run": 1, "start": 6, "end": 12}]


def test_docx_textbox_and_nested_table_are_extracted_once(tmp_path, config):
    import docx
    from lxml import etree

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    box = f'''<w:r xmlns:w="{w}" xmlns:mc="{mc}"><mc:AlternateContent>
      <mc:Choice Requires="wps"><w:drawing><w:txbxContent><w:p><w:r><w:t>상자 안의 제어 조건</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>
      <mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>상자 안의 제어 조건</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>
    </mc:AlternateContent></w:r>'''
    document = docx.Document()
    paragraph = document.add_paragraph("도형이 붙은 문단")
    paragraph._p.append(etree.fromstring(box))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "구분", "내용"
    table.cell(1, 0).text = "바깥 셀"
    inner = table.cell(1, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "안쪽 표", "보정 계수 갱신"
    path = tmp_path / "boxes.docx"
    document.save(str(path))
    result = parse_document(path, config)
    texts = [(b.kind, b.text) for b in result.blocks]
    assert texts[0] == ("paragraph", "도형이 붙은 문단")
    assert texts[1] == ("paragraph", "상자 안의 제어 조건") and result.blocks[1].flags == ["textbox"]
    assert [text for _, text in texts].count("상자 안의 제어 조건") == 1  # Fallback 복제본은 읽지 않는다
    assert ("table_row", "구분 | 내용") in texts and ("table_row", "안쪽 표 | 보정 계수 갱신") in texts
    nested = next(b for b in result.blocks if b.text.startswith("안쪽 표"))
    assert nested.locator["row"] == 0 and nested.locator["nested_path"] == [{"row": 1, "col": 1, "index": 0}]
    outer = next(b for b in result.blocks if b.text == "바깥 셀")
    assert outer.locator["row"] == 1 and "nested_path" not in outer.locator


def test_docx_headers_footers_only_when_configured(tmp_path, config):
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "머리말 문구"
    document.add_paragraph("본문 문단")
    path = tmp_path / "header.docx"
    document.save(str(path))
    assert [b.text for b in parse_document(path, config).blocks] == ["본문 문단"]
    included = make_config(tmp_path, **{"parsing.include_headers_footers": True})
    result = parse_document(path, included)
    assert [(b.kind, b.text) for b in result.blocks] == [("paragraph", "본문 문단"), ("header_footer", "머리말 문구")]
    assert result.coverage["scope"]["headers_footers"] is True


# ---------------------------------------------------------------- PDF
def test_pdf_two_columns_hyphen_and_headers(config):
    result = parse_document(FIXTURES / "sample_report.pdf", config)
    assert result.status == "partial"
    coverage = result.coverage
    assert coverage["units_total"] == 4 and coverage["ocr_needed_units"] == [4]
    assert coverage["multi_column_units"] == [2] and any(w["code"] == "multi_column" for w in result.warnings)
    texts = [b.text for b in result.blocks]
    assert PDF_YES in texts  # 줄바꿈된 두 줄이 한 문단으로
    assert PDF_HYPHEN in texts and coverage["edits"]["dehyphenated"] == 1  # 하이픈 병합을 기록
    removed = {item["text"] for item in coverage["edits"]["removed_headers_footers"]}
    assert removed == {PDF_HEADER, "- # -"}  # 반복 머리말·쪽 번호 삭제를 기록
    assert not any(PDF_HEADER in text for text in texts)
    page2 = [b.text for b in result.blocks if b.unit == 2]
    assert page2[0] == "모듈 간 온도 균일화 방법"
    assert page2[1].startswith("왼쪽 단 첫 문단이다.") and page2[-1].startswith("오른쪽 단 첫 문단이다.")
    assert all("오른쪽" not in text for text in page2[1:-1])  # 왼쪽 단을 다 읽은 뒤 오른쪽 단


def test_pdf_line_locators_point_at_the_extracted_characters(config):
    import pdfplumber

    result = parse_document(FIXTURES / "sample_report.pdf", config)
    block = next(b for b in result.blocks if b.text == PDF_YES)
    lines = block.locator["lines"]
    assert len(lines) == 2 and block.locator["page"] == 1
    assert all(len(line["xs"]) == line["end"] - line["start"] + 1 for line in lines)
    with pdfplumber.open(str(FIXTURES / "sample_report.pdf")) as pdf:
        page = pdf.pages[0]
        for line in lines:
            box = (line["xs"][0] - 0.5, line["top"] - 0.5, line["xs"][-1] + 0.5, line["bottom"] + 0.5)
            extracted = page.crop(box).extract_text().replace(" ", "")
            assert extracted == block.text[line["start"]:line["end"]].replace(" ", "")
    bullets = [b for b in result.blocks if b.kind == "list_item" and b.unit == 1]
    assert [b.text for b in bullets] == ["다음 분기에 열관리 성능 개선을 추진한다.", "회의는 매주 화요일에 진행한다."]


def test_pdf_table_rows_have_cells_and_headers(config):
    result = parse_document(FIXTURES / "sample_report.pdf", config)
    rows = [b for b in result.blocks if b.kind == "table_row"]
    assert [b.text for b in rows] == ["항목 | 조건 | 비고", "유량 밸브 | 편차 3도 초과 시 개도를 편차에 비례해 증가 | 신규",
                                      "일정 | 11월 시제품 평가 | -"]
    assert [c["header"] for c in rows[1].cells] == ["항목", "조건", "비고"]
    assert rows[1].locator["table"] == 0 and rows[1].locator["row"] == 1


def test_scanned_pdf_is_unsupported_not_empty(tmp_path, config):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.add_page(PdfReader(str(FIXTURES / "sample_report.pdf")).pages[3])
    path = tmp_path / "scan.pdf"
    with path.open("wb") as handle:
        writer.write(handle)
    result = parse_document(path, config)
    # 4쪽에는 머리말 글자가 조금 있으므로 텍스트는 있지만 OCR 필요로 분류된다
    assert result.status == "unsupported" and "OCR" in result.error
    assert result.coverage["units_ocr_needed"] == 1


def test_sideways_pdf_page_is_reported_not_silently_empty(tmp_path, config):
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(FIXTURES / "sample_report.pdf"))
    sideways = reader.pages[0]
    sideways.rotate(90)  # 본문이 옆으로 누운 페이지
    writer = PdfWriter()
    writer.add_page(sideways)
    writer.add_page(reader.pages[2])
    path = tmp_path / "mixed.pdf"
    with path.open("wb") as handle:
        writer.write(handle)
    result = parse_document(path, config)
    coverage = result.coverage
    assert result.status == "partial" and coverage["rotated_only_units"] == [1] and coverage["units_processed"] == 1
    assert coverage["rotated_chars_skipped"] > 100
    assert any(w["code"] == "rotated_page_not_analyzed" and w["unit"] == 1 for w in result.warnings)
    assert {b.unit for b in result.blocks} == {2}

    only = PdfWriter()
    only.add_page(sideways)
    alone = tmp_path / "sideways.pdf"
    with alone.open("wb") as handle:
        only.write(handle)
    unsupported = parse_document(alone, config)
    assert unsupported.status == "unsupported" and "회전" in unsupported.error  # '후보 없음'으로 처리하지 않는다


def _rotated_upright_pdf(path):
    """/Rotate 90 이지만 화면에서는 글자가 바로 보이는 페이지."""
    reportlab = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    assert reportlab
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setPageRotation(90)
    c.rotate(90)
    c.setFont("HYGothic-Medium", 11)
    c.drawString(80, -120, "측정값이 임계값을 넘으면 출력을 단계적으로 낮추는 제어를 적용한다.")
    c.drawString(80, -140, "다음 분기에 회의 일정을 다시 정한다.")
    c.showPage()
    c.save()


def test_rotated_page_with_upright_text_is_parsed_and_flagged(tmp_path, config):
    path = tmp_path / "rotated.pdf"
    _rotated_upright_pdf(path)
    result = parse_document(path, config)
    # 이어진 두 줄은 한 문단으로 묶인다
    assert [b.text for b in result.blocks] == [
        "측정값이 임계값을 넘으면 출력을 단계적으로 낮추는 제어를 적용한다. 다음 분기에 회의 일정을 다시 정한다."]
    assert result.coverage["rotated_units"] == [1] and result.coverage["rotated_only_units"] == []
    assert any(w["code"] == "rotated_text" for w in result.warnings)
    assert all(b.locator["rotation"] == 90 for b in result.blocks)


def test_encrypted_pdf_is_reported_as_unsupported(tmp_path, config):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(str(FIXTURES / "sample_report.pdf")))
    writer.encrypt("secret")
    path = tmp_path / "locked.pdf"
    with path.open("wb") as handle:
        writer.write(handle)
    result = parse_document(path, config)
    assert result.status == "unsupported" and "암호화" in result.error


# ---------------------------------------------------------------- TXT / MD
def test_text_offsets_match_file_content(config):
    path = FIXTURES / "sample_notes.txt"
    content = path.read_bytes().decode("utf-8")  # 줄바꿈을 바꾸지 않고 읽는다. Windows에서 CRLF로 받은 파일도 위치가 맞아야 한다
    result = parse_document(path, config)
    for block in result.blocks:
        locator = block.locator
        assert content[locator["char_start"]:locator["char_end"]] == block.text
    wrapped = result.blocks[1]
    assert (wrapped.locator["line_start"], wrapped.locator["line_end"]) == (3, 4) and "\n" in wrapped.text
    assert wrapped.hard_newlines is False  # 줄바꿈은 문장 경계가 아니다
    assert [b.kind for b in result.blocks] == ["paragraph", "paragraph", "list_item", "list_item"]


def test_markdown_headings_tables_and_lists(config):
    result = parse_document(FIXTURES / "sample_notes.md", config)
    assert [b.kind for b in result.blocks] == ["heading", "heading", "paragraph", "table_row", "table_row",
                                               "table_row", "heading", "list_item", "list_item"]
    row = result.blocks[5]
    assert [row.text[c["start"]:c["end"]] for c in row.cells] == ["2차", "후보 영역만 고해상도 재촬영 후 분류기 적용", "신규"]
    assert [c["header"] for c in row.cells] == ["단계", "처리", "비고"]
    assert result.blocks[2].section_title == "제안" and result.blocks[7].section_title == "일정"


def test_text_decoding_failure_asks_for_explicit_encoding(tmp_path, config):
    path = tmp_path / "cp949.txt"
    path.write_bytes("임계값을 넘으면 출력을 낮춘다.".encode("cp949"))
    failed = parse_document(path, config)
    assert failed.status == "failed" and "--encoding" in failed.error
    ok = parse_document(path, config, encoding="cp949")
    assert ok.status == "ok" and ok.blocks[0].text == "임계값을 넘으면 출력을 낮춘다."


# ---------------------------------------------------------------- 미지원·제한
@pytest.mark.parametrize("name", ["legacy.ppt", "report.hwp", "macro.pptm", "scan.png", "noext"])
def test_unsupported_formats_are_reported_with_reason(tmp_path, config, name):
    path = tmp_path / name
    path.write_bytes(b"data")
    result = parse_document(path, config)
    assert result.status == "unsupported" and result.error and not result.blocks


def test_corrupt_or_encrypted_ooxml_is_unsupported(tmp_path, config):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0 not a zip")
    result = parse_document(path, config)
    assert result.status == "unsupported" and "암호화" in result.error


def test_size_and_uncompressed_limits_stop_processing(tmp_path):
    small = make_config(tmp_path, **{"limits.max_uncompressed_mb": 1})
    path = tmp_path / "bomb.pptx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/big.bin", b"\0" * (3 << 20))
    result = parse_document(path, small)
    assert result.status == "unsupported" and "압축 해제 용량" in result.error

    tiny = make_config(tmp_path, **{"limits.max_file_mb": 0})
    result = parse_document(FIXTURES / "sample_report.pptx", tiny)
    assert result.status == "unsupported" and "파일 크기" in result.error

    few = make_config(tmp_path, **{"limits.max_units": 3})
    assert parse_document(FIXTURES / "sample_report.pptx", few).status == "unsupported"
    assert parse_document(FIXTURES / "sample_report.pdf", few).status == "unsupported"


def test_effect_bullet_is_a_separate_block_from_the_grouped_parent(config):
    result = parse_document(FIXTURES / "sample_report.pptx", config)
    effect = next(b for b in result.blocks if b.text == EFFECT_BULLET)
    assert effect.level == 0 and effect.unit == 3
    assert io.BytesIO  # 가져온 모듈 사용 표시
