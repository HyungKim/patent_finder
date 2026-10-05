"""PPTX 파서. 슬라이드별 텍스트 상자·표·그룹 도형·SmartArt·차트 제목을 추출한다 (스펙 4.1, 4.2).

XML 순서는 사람이 읽는 순서와 다를 수 있으므로 좌표 기반 읽기 순서를 계산하고,
두 순서를 모두 locator에 기록한다.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from ..config import AppConfig
from ..existing_marks import clear_pptx_marks, describe_marks, total_marks
from .base import Block, ParseResult, UnsupportedDocument, check_zip_limits, finalize_status

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
OLE_URI = "http://schemas.openxmlformats.org/presentationml/2006/ole"

Transform = Callable[[float, float, float, float], tuple[float, float, float, float]]


def a_paragraph_text(p: Any) -> tuple[str, list[dict[str, int]]]:
    """a:p의 텍스트와 run별 문자 구간. run 번호는 a:p의 자식 순번이다."""
    parts: list[str] = []
    runs: list[dict[str, int]] = []
    position = 0
    for index, child in enumerate(p):
        tag = child.tag
        if tag in (A + "r", A + "fld"):
            node = child.find(A + "t")
            piece = (node.text or "") if node is not None else ""
        elif tag == A + "br":
            piece = "\n"
        else:
            continue
        if piece:
            runs.append({"run": index, "start": position, "end": position + len(piece)})
            parts.append(piece)
            position += len(piece)
    return "".join(parts), runs


def _identity(x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
    return x, y, w, h


def group_transform(group_element: Any, parent: Transform) -> Transform:
    """그룹 도형의 자식 좌표를 슬라이드 좌표로 바꾸는 변환 (회전·뒤집기는 반영하지 않음)."""
    xfrm = None
    group_props = group_element.find(P + "grpSpPr")
    if group_props is not None:
        xfrm = group_props.find(A + "xfrm")
    if xfrm is None:
        return parent
    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
    ch_off, ch_ext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return parent
    ox, oy = float(off.get("x", 0)), float(off.get("y", 0))
    cx, cy = float(ext.get("cx", 0)), float(ext.get("cy", 0))
    cox, coy = float(ch_off.get("x", 0)), float(ch_off.get("y", 0))
    ccx, ccy = float(ch_ext.get("cx", 0)), float(ch_ext.get("cy", 0))
    sx = cx / ccx if ccx else 1.0
    sy = cy / ccy if ccy else 1.0

    def transform(x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
        return parent(ox + (x - cox) * sx, oy + (y - coy) * sy, w * sx, h * sy)

    return transform


@dataclass
class ShapeInfo:
    shape: Any
    left: float
    top: float
    width: float
    height: float
    xml_order: int
    group_path: list[int] = field(default_factory=list)
    reading_order: int = 0

    @property
    def bbox(self) -> dict[str, int]:
        return {"left": int(self.left), "top": int(self.top), "width": int(self.width), "height": int(self.height)}


def collect_shapes(shapes: Any, transform: Transform = _identity,
                   group_path: list[int] | None = None, out: list[ShapeInfo] | None = None) -> list[ShapeInfo]:
    """그룹을 풀어 모든 말단 도형을 슬라이드 절대 좌표와 함께 모은다."""
    from pptx.shapes.group import GroupShape

    out = [] if out is None else out
    group_path = group_path or []
    for shape in shapes:
        try:
            box = transform(float(shape.left or 0), float(shape.top or 0),
                            float(shape.width or 0), float(shape.height or 0))
        except Exception:
            box = (0.0, 0.0, 0.0, 0.0)
        if isinstance(shape, GroupShape):
            collect_shapes(shape.shapes, group_transform(shape._element, transform),
                           group_path + [shape.shape_id], out)
        else:
            out.append(ShapeInfo(shape, *box, xml_order=len(out), group_path=list(group_path)))
    return out


def assign_reading_order(infos: list[ShapeInfo]) -> list[ShapeInfo]:
    """위→아래, 같은 행에서는 왼쪽→오른쪽. 세로로 절반 이상 겹치면 같은 행으로 본다."""
    rows: list[list[ShapeInfo]] = []
    for info in sorted(infos, key=lambda i: (i.top, i.left)):
        placed = False
        for row in rows:
            anchor = row[0]
            overlap = min(anchor.top + anchor.height, info.top + info.height) - max(anchor.top, info.top)
            base = min(anchor.height, info.height)
            if base > 0 and overlap / base >= 0.5:
                row.append(info)
                placed = True
                break
        if not placed:
            rows.append([info])
    ordered: list[ShapeInfo] = []
    for row in rows:
        ordered.extend(sorted(row, key=lambda i: (i.left, i.top)))
    for index, info in enumerate(ordered):
        info.reading_order = index
    return ordered


def _is_title(shape: Any) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        from pptx.enum.shapes import PP_PLACEHOLDER

        return shape.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE,
        )
    except Exception:
        return False


def _graphic_uri(shape: Any) -> str | None:
    data = shape._element.find(".//" + A + "graphicData")
    return None if data is None else data.get("uri")


class _SlideParser:
    def __init__(self, slide: Any, number: int, config: AppConfig, slide_size: tuple[int, int]) -> None:
        self.slide = slide
        self.number = number
        self.config = config
        self.slide_size = slide_size
        self.blocks: list[Block] = []
        self.stats = {"pictures": 0, "picture_area": 0.0, "ole_objects": 0, "charts": 0, "smartart": 0}

    def _locator(self, info: ShapeInfo, **extra: Any) -> dict[str, Any]:
        locator: dict[str, Any] = {
            "format": "pptx", "slide": self.number, "shape_id": info.shape.shape_id,
            "shape_name": info.shape.name, "bbox_emu": info.bbox,
            "xml_order": info.xml_order, "reading_order": info.reading_order,
        }
        if info.group_path:
            locator["group_path"] = info.group_path
        locator.update(extra)
        return locator

    def _add(self, kind: str, text: str, locator: dict[str, Any], **extra: Any) -> None:
        if text.strip():
            self.blocks.append(Block(kind=kind, text=text, locator=locator, unit=self.number,
                                     section_key=f"slide:{self.number}", **extra))

    def parse(self) -> list[Block]:
        infos = assign_reading_order(collect_shapes(self.slide.shapes))
        titles = [info for info in infos if _is_title(info.shape)]
        for info in titles + [info for info in infos if info not in titles]:
            self._shape(info, is_title=info in titles)
        self._infer_title()
        title = next((block.text.strip() for block in self.blocks if block.kind == "title"), None)
        for block in self.blocks:
            block.section_title = title
        if self.config.parsing.include_speaker_notes:
            self._notes(title)
        return self.blocks

    def _shape(self, info: ShapeInfo, is_title: bool) -> None:
        shape = info.shape
        uri = _graphic_uri(shape)
        if getattr(shape, "has_table", False):
            self._table(info)
        elif getattr(shape, "has_chart", False):
            self._chart(info)
        elif uri == DIAGRAM_URI:
            self._smartart(info)
        elif uri == OLE_URI:
            self.stats["ole_objects"] += 1
        elif shape._element.tag == P + "pic":
            self.stats["pictures"] += 1
            area = self.slide_size[0] * self.slide_size[1]
            if area:
                self.stats["picture_area"] += (info.width * info.height) / area
        elif getattr(shape, "has_text_frame", False):
            body = shape._element.find(P + "txBody")
            if body is not None:
                self._text_body(info, body, is_title)

    def _text_body(self, info: ShapeInfo, body: Any, is_title: bool) -> None:
        items = []
        for index, p in enumerate(body.findall(A + "p")):
            text, runs = a_paragraph_text(p)
            if not text.strip():
                continue
            ppr = p.find(A + "pPr")
            level_value = ppr.get("lvl") if ppr is not None else None
            level = int(level_value) if level_value and level_value.isdigit() else 0
            items.append((index, text, runs, level))
        if is_title:
            kind = "title"
        elif len(items) >= 2 or any(level > 0 for _, _, _, level in items):
            kind = "list_item"
        else:
            kind = "paragraph"
        container = f"s{self.number}:sh{info.shape.shape_id}"
        for index, text, runs, level in items:
            self._add(kind, text, self._locator(info, paragraph=index, runs=runs),
                      level=level, container=container)

    def _table(self, info: ShapeInfo) -> None:
        tbl = info.shape._element.find(".//" + A + "tbl")
        if tbl is None:
            return
        props = tbl.find(A + "tblPr")
        first_row = props is not None and props.get("firstRow") in ("1", "true")
        rows = tbl.findall(A + "tr")
        headers: dict[int, str] = {}
        for row_index, tr in enumerate(rows):
            pieces: list[str] = []
            cells: list[dict[str, Any]] = []
            position = 0
            is_header = first_row and row_index == 0 and len(rows) >= 2
            for col, tc in enumerate(tr.findall(A + "tc")):
                if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
                    continue
                body = tc.find(A + "txBody")
                paragraphs = []
                if body is not None:
                    for para_index, p in enumerate(body.findall(A + "p")):
                        text, runs = a_paragraph_text(p)
                        if text.strip():
                            paragraphs.append((para_index, text, runs))
                cell_text = "\n".join(text for _, text, _ in paragraphs)
                if not cell_text.strip():
                    continue
                if is_header:
                    headers[col] = " ".join(cell_text.split())
                if pieces:
                    pieces.append(" | ")
                    position += 3
                para_locs = []
                offset = position
                for para_index, text, runs in paragraphs:
                    para_locs.append({"index": para_index, "start": offset, "end": offset + len(text), "runs": runs})
                    offset += len(text) + 1
                cells.append({"col": col, "start": position, "end": position + len(cell_text),
                              "header": None if is_header else headers.get(col), "paragraphs": para_locs})
                pieces.append(cell_text)
                position += len(cell_text)
            if pieces:
                self._add("table_row", "".join(pieces), self._locator(info, table_row=row_index),
                          cells=cells, flags=["table_header"] if is_header else [])

    def _chart(self, info: ShapeInfo) -> None:
        self.stats["charts"] += 1
        try:
            chart = info.shape.chart
            if chart.has_title and chart.chart_title.has_text_frame:
                text = chart.chart_title.text_frame.text
                self._add("diagram_text", text, self._locator(info, chart="title"), flags=["chart_title"])
        except Exception:
            return

    def _smartart(self, info: ShapeInfo) -> None:
        from lxml import etree

        self.stats["smartart"] += 1
        rel_ids = info.shape._element.find(".//" + DGM + "relIds")
        rid = rel_ids.get(R + "dm") if rel_ids is not None else None
        if not rid:
            return
        try:
            part = self.slide.part.related_part(rid)
            root = etree.fromstring(part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
        except Exception:
            return
        container = f"s{self.number}:sh{info.shape.shape_id}"
        for point in root.iter(DGM + "pt"):
            if point.get("type") in ("pres", "parTrans", "sibTrans"):
                continue
            body = point.find(DGM + "t")
            if body is None:
                continue
            for index, p in enumerate(body.findall(A + "p")):
                text, _runs = a_paragraph_text(p)
                self._add("diagram_text", text,
                          self._locator(info, smartart_point=point.get("modelId"), paragraph=index),
                          flags=["smartart"], container=container)

    def _infer_title(self) -> None:
        if any(block.kind == "title" for block in self.blocks) or not self.blocks:
            return
        first = min(self.blocks, key=lambda b: b.locator.get("reading_order", 0))
        top = first.locator.get("bbox_emu", {}).get("top", 0)
        same_shape = [b for b in self.blocks if b.locator.get("shape_id") == first.locator.get("shape_id")]
        if (first.kind in ("paragraph", "list_item") and len(same_shape) == 1
                and len(first.text.strip()) <= 80 and self.slide_size[1]
                and top < 0.2 * self.slide_size[1]):
            first.kind = "title"
            first.flags.append("title_inferred")
            self.blocks.remove(first)
            self.blocks.insert(0, first)

    def _notes(self, title: str | None) -> None:
        if not self.slide.has_notes_slide:
            return
        frame = self.slide.notes_slide.notes_text_frame
        if frame is None:
            return
        for index, p in enumerate(frame._txBody.findall(A + "p")):
            text, runs = a_paragraph_text(p)
            if text.strip():
                self.blocks.append(Block(
                    kind="notes", text=text, unit=self.number,
                    locator={"format": "pptx", "slide": self.number, "notes": True, "paragraph": index, "runs": runs},
                    section_key=f"slide:{self.number}", section_title=title, flags=["speaker_notes"],
                ))


def parse_pptx(path: Path, config: AppConfig) -> ParseResult:
    from pptx import Presentation

    result = ParseResult(format="pptx", status="ok")
    check_zip_limits(path, config.limits.max_uncompressed_mb)
    presentation = Presentation(str(path))
    # 기존 마킹을 걷어낸 구조로 파싱한다. 메모리 안에서만 바꾸며 파일은 저장하지 않는다.
    existing = clear_pptx_marks(presentation)
    slides = list(presentation.slides)
    if len(slides) > config.limits.max_units:
        raise UnsupportedDocument(
            "unit_limit", f"슬라이드 수({len(slides)})가 제한({config.limits.max_units})을 넘습니다. limits.max_units를 확인하세요."
        )
    slide_size = (int(presentation.slide_width or 0), int(presentation.slide_height or 0))
    totals = {"pictures": 0, "ole_objects": 0, "charts": 0, "smartart": 0}
    failed: list[int] = []
    hidden: list[int] = []
    image_only: list[int] = []
    processed = 0
    for number, slide in enumerate(slides, start=1):
        is_hidden = slide._element.get("show") in ("0", "false")
        if is_hidden:
            hidden.append(number)
            if not config.parsing.include_hidden_slides:
                continue
        parser = _SlideParser(slide, number, config, slide_size)
        try:
            blocks = parser.parse()
        except Exception as exc:  # 한 슬라이드의 오류가 전체를 막지 않게 한다
            failed.append(number)
            result.warn("slide_failed", f"슬라이드 {number} 처리 실패: {type(exc).__name__}", unit=number)
            continue
        processed += 1
        if is_hidden:
            for block in blocks:
                block.flags.append("hidden_slide")
        for key in totals:
            totals[key] += parser.stats[key]
        body_chars = sum(len(b.text.strip()) for b in blocks if b.kind != "title")
        if parser.stats["pictures"] and body_chars < 20 and parser.stats["picture_area"] >= 0.3:
            image_only.append(number)
            result.warn("image_only_slide",
                        f"슬라이드 {number}: 이미지 위주이며 텍스트가 거의 없습니다. 이미지 안의 내용은 분석되지 않았습니다 (OCR 미지원).",
                        unit=number)
        result.blocks.extend(blocks)

    result.coverage = {
        "unit_type": "slide",
        "units_total": len(slides),
        "units_processed": processed,
        "units_failed": len(failed),
        "failed_units": failed,
        "units_ocr_needed": len(image_only),
        "ocr_needed_units": image_only,
        "hidden_units": hidden,
        "blocks": len(result.blocks),
        "slide_size_emu": list(slide_size),
        "scope": {
            "speaker_notes": config.parsing.include_speaker_notes,
            "hidden_slides": config.parsing.include_hidden_slides,
        },
        "not_analyzed": {"pictures": totals["pictures"], "ole_objects": totals["ole_objects"],
                         "chart_data": totals["charts"]},
        "smartart_shapes": totals["smartart"],
        "existing_marks": existing,
    }
    if total_marks(existing):
        result.warn("existing_marks_cleared",
                    f"기존 마킹 {total_marks(existing)}건({describe_marks(existing)})을 걷어내고 분석했습니다. "
                    "마킹 사본에도 남기지 않습니다.")
    if totals["pictures"]:
        result.warn("images_not_analyzed", f"이미지 {totals['pictures']}개 안의 텍스트는 분석하지 않았습니다 (OCR 미지원).")
    if totals["charts"]:
        result.warn("chart_data_not_analyzed", f"차트 {totals['charts']}개는 제목만 추출했습니다 (데이터·범례 해석 제외).")
    if totals["ole_objects"]:
        result.warn("objects_not_analyzed", f"포함 개체 {totals['ole_objects']}개는 분석하지 않았습니다.")
    return finalize_status(result)
