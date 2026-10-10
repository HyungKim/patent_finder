"""원문 사본 마킹: PPTX, PDF (스펙 13절, 변경안 A4).

- 원본은 절대 수정하지 않는다. 항상 별도 파일에 쓰고, 끝난 뒤 원본 해시가 그대로인지 확인한다.
- PPTX: 후보 문단의 run에 강조색을 넣고(run 단위), 도형 둘레에 테두리와 후보 꼬리표를 붙인다.
- PDF: 문자 좌표로 만든 quad로 highlight 주석을 단다.
- 원본에 이미 있던 마킹(도구 표시, 형광펜, 메모, PDF 표시 주석)은 먼저 걷어내고 새로 표시한다.
- 위치를 확신할 수 없으면(회전 페이지, 좌표 없는 도형 등) 칠하지 않고 사유를 남긴다.
- DOCX 사본 마킹은 아직 구현하지 않았다(HTML 보고서의 위치 목록을 사용).
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from ..existing_marks import TOOL_AUTHOR, clear_pdf_marks, clear_pptx_marks
from ..runtime import sha256_file

HIGHLIGHT_RGB = "FFE066"
OUTLINE_RGB = "D97706"
TAG_RGB = "9A4A06"
_SUMMARY_LINES = 12


def _tag_text(numbers: list[int], score: float | None, system: bool, stage: str | None) -> str:
    ids = ", ".join(f"#{n}" for n in numbers)
    if not system:
        return f"검토자 YES {ids}"
    prefix = "임시 후보(SEED)" if stage == "SEED" else "특허검토 후보"
    return f"{prefix} {ids}" + (f" · {score:.2f}" if score is not None else "")


def _relative_spans(spans: list[tuple[int, int]], start: int, end: int) -> list[tuple[int, int]]:
    """문단 기준 구간 중 [start, end) 조각과 겹치는 부분을 조각 기준 offset으로 바꾼다."""
    result = []
    for a, b in spans:
        lo, hi = max(a, start), min(b, end)
        if hi > lo:
            result.append((lo - start, hi - start))
    return result


def marked_paragraphs(document: dict[str, Any]) -> list[dict[str, Any]]:
    """표시 대상 문단에 문서 순서대로 번호를 붙인다."""
    items = []
    for paragraph in document["paragraphs"]:
        if not paragraph["marked_spans"]:
            continue
        marked = [segment for segment in paragraph["segments"] if segment["marked"]]
        scores = [s["stage1_score"] for s in marked if s["stage1_score"] is not None and s["final_decision"] == "CANDIDATE"]
        items.append({
            "number": len(items) + 1, "paragraph": paragraph, "spans": paragraph["marked_spans"],
            "score": max(scores) if scores else None,
            "system": any(s["final_decision"] == "CANDIDATE" for s in marked),
        })
    return items


# ====================================================================== PPTX
def _find_shape(shapes: Any, shape_id: int) -> Any:
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if shape.shape_id == shape_id:
            return shape
        if isinstance(shape, GroupShape):
            found = _find_shape(shape.shapes, shape_id)
            if found is not None:
                return found
    return None


def _highlight_run(run: Any) -> None:
    """a:r / a:fld에 강조색을 넣는다. a:rPr 안의 요소 순서(스키마)를 지킨다."""
    from lxml import etree

    from ..parsers.pptx import A

    properties = run.find(A + "rPr")
    if properties is None:
        properties = etree.Element(A + "rPr")
        run.insert(0, properties)
    for old in properties.findall(A + "highlight"):
        properties.remove(old)
    before = {A + name for name in ("ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
                                    "effectLst", "effectDag")}
    position = 0
    for index, child in enumerate(properties):
        if child.tag in before:
            position = index + 1
    highlight = etree.Element(A + "highlight")
    etree.SubElement(highlight, A + "srgbClr", val=HIGHLIGHT_RGB)
    properties.insert(position, highlight)


def _split_run(p: Any, element: Any, run: dict[str, int], covered: list[tuple[int, int]]) -> int:
    """run의 일부만 표시 대상일 때 run을 조각으로 나누고 해당 조각만 칠한다. 서식(a:rPr)은 조각마다 복사한다."""
    import copy

    from ..parsers.pptx import A

    text = element.find(A + "t").text or ""
    cuts = sorted({run["start"], run["end"], *(a for a, _ in covered), *(b for _, b in covered)})
    pieces = []
    for a, b in zip(cuts, cuts[1:]):
        if b > a:
            pieces.append((text[a - run["start"]:b - run["start"]], any(ca <= a and b <= cb for ca, cb in covered)))
    position = list(p).index(element)
    p.remove(element)
    count = 0
    for offset, (piece, marked) in enumerate(pieces):
        clone = copy.deepcopy(element)
        clone.find(A + "t").text = piece
        if marked:
            _highlight_run(clone)
            count += 1
        p.insert(position + offset, clone)
    return count


def _highlight_paragraph(p: Any, runs: list[dict[str, int]], spans: list[tuple[int, int]]) -> int:
    from ..parsers.pptx import A

    children = list(p)
    count = 0
    # 뒤 run부터 처리한다: 앞 run을 조각으로 나눠도 아직 처리하지 않은 run의 번호가 밀리지 않는다
    for run in sorted(runs, key=lambda item: item["run"], reverse=True):
        if run["run"] >= len(children):
            continue
        element = children[run["run"]]
        if not (element.tag.endswith("}r") or element.tag.endswith("}fld")):
            continue
        covered = [(max(a, run["start"]), min(b, run["end"])) for a, b in spans if run["start"] < b and a < run["end"]]
        if not covered:
            continue
        whole = len(covered) == 1 and covered[0] == (run["start"], run["end"])
        if whole or element.tag.endswith("}fld") or element.find(A + "t") is None:
            _highlight_run(element)
            count += 1
        else:
            count += _split_run(p, element, run, covered)
    return count


def annotate_pptx(source: Path, destination: Path, items: list[dict[str, Any]], stage: str | None,
                  summary_slide: bool = True) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    from ..parsers.pptx import A, a_paragraph_text

    presentation = Presentation(str(source))
    cleared = clear_pptx_marks(presentation)  # 파서와 같은 함수: 같은 슬라이드 번호·도형 구조에서 시작
    slides = list(presentation.slides)
    report: dict[str, Any] = {"format": "pptx", "marked": 0, "skipped": [], "cleared": cleared,
                              "granularity": "run 강조 + 도형 테두리"}
    boxes: dict[tuple[int, int], dict[str, Any]] = {}
    corner_tags: dict[int, list[str]] = {}

    for item in items:
        paragraph = item["paragraph"]
        done = False
        for part in paragraph["locator"]["parts"]:
            relative = _relative_spans(item["spans"], part["start"], part["end"])
            locator = part["locator"]
            if not relative:
                continue
            slide_number = locator.get("slide")
            if not slide_number or slide_number > len(slides):
                continue
            slide = slides[slide_number - 1]
            if locator.get("notes"):
                corner_tags.setdefault(slide_number, []).append(
                    _tag_text([item["number"]], item["score"], item["system"], stage) + " (발표자 노트)")
                done = True
                continue
            shape = _find_shape(slide.shapes, locator.get("shape_id"))
            if shape is None:
                continue
            if "table_row" in locator:
                table = shape._element.find(".//" + A + "tbl")
                rows = table.findall(A + "tr") if table is not None else []
                if locator["table_row"] < len(rows):
                    cells = rows[locator["table_row"]].findall(A + "tc")
                    for cell in paragraph["locator"].get("cells", []):
                        if cell["col"] < len(cells) and any(cell["start"] < b and a < cell["end"] for a, b in item["spans"]):
                            for p in cells[cell["col"]].iter(A + "p"):
                                _text, runs = a_paragraph_text(p)
                                _highlight_paragraph(p, runs, [(0, 1 << 30)])
            elif "paragraph" in locator and "smartart_point" not in locator:
                body = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
                paragraphs = body.findall(A + "p") if body is not None else []
                if locator["paragraph"] < len(paragraphs):
                    _highlight_paragraph(paragraphs[locator["paragraph"]], locator.get("runs", []), relative)
            box = locator.get("bbox_emu") or {}
            key = (slide_number, locator.get("shape_id"))
            entry = boxes.setdefault(key, {"bbox": box, "numbers": [], "score": None, "system": False})
            if item["number"] not in entry["numbers"]:
                entry["numbers"].append(item["number"])
            if item["score"] is not None:
                entry["score"] = max(entry["score"] or 0.0, item["score"])
            entry["system"] = entry["system"] or item["system"]
            done = True
        if done:
            report["marked"] += 1
        else:
            report["skipped"].append({"number": item["number"], "reason": "원문 도형을 찾지 못해 표시하지 않음"})

    tag_height = Inches(0.3)
    for (slide_number, _shape_id), entry in boxes.items():
        slide = slides[slide_number - 1]
        box = entry["bbox"]
        text = _tag_text(entry["numbers"], entry["score"], entry["system"], stage)
        if box.get("width") and box.get("height"):
            outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(box["left"]), Emu(box["top"]),
                                             Emu(box["width"]), Emu(box["height"]))
            outline.name = "PM_MARK_OUTLINE"
            outline.fill.background()
            outline.line.color.rgb = RGBColor.from_string(OUTLINE_RGB)
            outline.line.width = Pt(2.25)
            outline.shadow.inherit = False
            left, top = box["left"], max(0, box["top"] - int(tag_height))
        else:
            left, top = 0, 0
        _add_tag(slide, left, top, text, Inches, Pt, RGBColor, Emu)
    for slide_number, texts in corner_tags.items():
        for row, text in enumerate(texts):
            _add_tag(slides[slide_number - 1], 0, int(tag_height) * row, text, Inches, Pt, RGBColor, Emu)

    if summary_slide and items:
        _add_summary_slides(presentation, items, stage, Inches, Pt)
    presentation.save(str(destination))
    return report


def _add_tag(slide: Any, left: int, top: int, text: str, Inches: Any, Pt: Any, RGBColor: Any, Emu: Any) -> None:
    width = Inches(min(6.0, 0.9 + 0.16 * len(text)))
    tag = slide.shapes.add_textbox(Emu(left), Emu(top), width, Inches(0.3))
    tag.name = "PM_MARK_TAG"
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor.from_string(TAG_RGB)
    frame = tag.text_frame
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.word_wrap = False
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _add_summary_slides(presentation: Any, items: list[dict[str, Any]], stage: str | None, Inches: Any, Pt: Any) -> None:
    layout = min(presentation.slide_layouts, key=lambda item: len(item.placeholders))
    width = presentation.slide_width
    notes = {
        "SEED": "임시 후보(SEED · 미검증): 합성 예문 기반 분류기의 결과이며 놓친 후보가 있을 수 있습니다.",
        "PILOT": "파일럿 모델의 결과입니다. 평가 보고서의 Recall을 확인하세요.",
    }
    for start in range(0, len(items), _SUMMARY_LINES):
        slide = presentation.slides.add_slide(layout)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), width - Inches(1.0), Inches(0.6))
        title.name = "PM_SUMMARY_TITLE"
        run = title.text_frame.paragraphs[0].add_run()
        run.text = "특허 검토 후보 요약 (자동 생성)"
        run.font.size = Pt(24)
        run.font.bold = True
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), width - Inches(1.0), Inches(5.5))
        body.name = "PM_SUMMARY_BODY"
        frame = body.text_frame
        frame.word_wrap = True
        lines = [notes.get(stage or "", "") or "표시는 검토 후보이며 등록 가능성·신규성·침해 여부를 판정하지 않습니다."]
        for item in items[start: start + _SUMMARY_LINES]:
            text = " ".join(item["paragraph"]["original_text"].split())
            text = text if len(text) <= 70 else text[:70] + "…"
            score = f" · 후보 점수 {item['score']:.2f}" if item["score"] is not None else " · 검토자 YES"
            lines.append(f"#{item['number']} · {item['paragraph']['location']}{score} · {text}")
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = line
            run.font.size = Pt(12 if index else 11)


# ====================================================================== PDF
def annotate_pdf(source: Path, destination: Path, items: list[dict[str, Any]], stage: str | None) -> dict[str, Any]:
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Highlight
    from pypdf.generic import ArrayObject, FloatObject, NameObject, TextStringObject

    reader = PdfReader(str(source))
    writer = PdfWriter(clone_from=reader)
    cleared = clear_pdf_marks(writer)
    report: dict[str, Any] = {"format": "pdf", "marked": 0, "skipped": [], "cleared": cleared,
                              "granularity": "문자 좌표 quad"}
    for item in items:
        paragraph = item["paragraph"]
        quads_by_page: dict[int, list[float]] = {}
        reason = None
        for part in paragraph["locator"]["parts"]:
            locator = part["locator"]
            relative = _relative_spans(item["spans"], part["start"], part["end"])
            page_number = locator.get("page")
            if not relative or not page_number or page_number > len(writer.pages):
                continue
            page = writer.pages[page_number - 1]
            if locator.get("rotation") or (page.rotation or 0) % 360:
                reason = "회전된 페이지라 위치를 확신할 수 없어 표시하지 않음"
                continue
            origin_x, origin_y = float(page.mediabox.left), float(page.mediabox.bottom)
            for line in locator.get("lines", []):
                for a, b in relative:
                    lo, hi = max(a, line["start"]), min(b, line["end"])
                    if hi <= lo:
                        continue
                    xs = line["xs"]
                    x0, x1 = xs[lo - line["start"]] + origin_x, xs[hi - line["start"]] + origin_x
                    y0, y1 = line["y0"] + origin_y, line["y1"] + origin_y
                    quads_by_page.setdefault(page_number, []).extend([x0, y1, x1, y1, x0, y0, x1, y0])
        if not quads_by_page:
            report["skipped"].append({"number": item["number"], "reason": reason or "문자 좌표가 없어 표시하지 않음"})
            continue
        text = _tag_text([item["number"]], item["score"], item["system"], stage)
        for page_number, quads in quads_by_page.items():
            xs, ys = quads[0::2], quads[1::2]
            annotation = Highlight(
                rect=(min(xs), min(ys), max(xs), max(ys)),
                quad_points=ArrayObject([FloatObject(round(value, 2)) for value in quads]),
                highlight_color=HIGHLIGHT_RGB,
            )
            annotation[NameObject("/Contents")] = TextStringObject(text)
            annotation[NameObject("/T")] = TextStringObject(TOOL_AUTHOR)
            writer.add_annotation(page_number=page_number - 1, annotation=annotation)
        report["marked"] += 1
    # 걷어낸 주석 개체가 파일에 남지 않게 참조되지 않는 개체를 정리한다.
    try:
        writer.compress_identical_objects(remove_duplicates=False, remove_unreferenced=True)
    except TypeError:  # 예전 pypdf의 인자 이름
        writer.compress_identical_objects(remove_identicals=False, remove_orphans=True)
    with destination.open("wb") as handle:
        writer.write(handle)
    return report


# ====================================================================== 공통
def annotate_documents(run: dict[str, Any], output_dir: Path, summary_slide: bool = True) -> list[dict[str, Any]]:
    """run의 PPTX/PDF 문서마다 표시된 사본을 만든다. 표시할 구간이 없는 문서도 사유를 보고한다."""
    output_dir = Path(output_dir)
    results = []
    used_names: set[str] = set()
    stage = run["model_stage"] if run["classifier_version"] else None
    for document in run["documents"]:
        if not document["document_id"] or document["run_status"] not in ("ok", "partial"):
            continue
        fmt = document["format"]
        entry: dict[str, Any] = {"file_name": document["file_name"], "format": fmt}
        if fmt not in ("pptx", "pdf"):
            entry.update(status="not_implemented",
                         reason="이 형식의 사본 마킹은 아직 구현되지 않았습니다. HTML 보고서의 위치 목록을 사용하세요.")
            results.append(entry)
            continue
        items = marked_paragraphs(document)
        source = Path(document["local_path"])
        if not items:
            entry.update(status="no_marks", reason="표시할 구간이 없습니다.")
        elif not source.is_file() or sha256_file(source) != document["content_sha256"]:
            entry.update(status="skipped", reason="원본 파일이 없거나 분석 이후 바뀌어 사본을 만들지 않았습니다.")
        else:
            output_dir.mkdir(parents=True, exist_ok=True)
            # 마킹 사본을 다시 넣은 경우 이름이 x.marked.marked 가 되지 않게 하고, 같은 이름이 겹치면 번호를 붙인다.
            stem = source.stem[:-len(".marked")] if source.stem.lower().endswith(".marked") else source.stem
            name, number = f"{stem}.marked{source.suffix.lower()}", 2
            while name.lower() in used_names:
                name, number = f"{stem}.marked-{number}{source.suffix.lower()}", number + 1
            used_names.add(name.lower())
            destination = output_dir / name
            if destination.resolve() == source.resolve():
                entry.update(status="skipped", reason="출력 경로가 원본과 같습니다.")
            else:
                temporary = destination.with_name(destination.name + ".part")
                try:
                    if fmt == "pptx":
                        report = annotate_pptx(source, temporary, items, stage, summary_slide)
                    else:
                        report = annotate_pdf(source, temporary, items, stage)
                    os.replace(temporary, destination)
                    entry.update(status="ok", output=str(destination), candidates=len(items), **report)
                except Exception as exc:  # 사본 생성 실패가 다른 결과물을 막지 않게 한다
                    temporary.unlink(missing_ok=True)
                    entry.update(status="failed", reason=f"{type(exc).__name__}: {exc}")
                if sha256_file(source) != document["content_sha256"]:
                    raise RuntimeError(f"원본 파일이 변경되었습니다: {source}")
        results.append(entry)
    return results
